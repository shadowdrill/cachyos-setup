"""Generate the KRDP Remote Desktop guide (Word) with python-docx."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pathlib import Path

OUT = str(Path(__file__).resolve().parents[2] / "docs" / "rdp" / "KRDP-Remote-Desktop-Guide.docx")
ACCENT = RGBColor(0x1D, 0x4E, 0x89)
MUTED = RGBColor(0x55, 0x5B, 0x66)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True


def shade(el, fill):
    props = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def rich(par, text):
    """Add text where `backticks` become monospace and **stars** become bold."""
    import re
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        else:
            par.add_run(part)
    return par


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, style=None):
    return rich(doc.add_paragraph(style=style), text)


def bullets(items, style="List Bullet"):
    for it in items:
        p(it, style)


_list_num = {"id": None}


def new_list():
    """Start a fresh 1,2,3 sequence: a new w:num on List Number's abstractNum with a restart."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num = doc.styles["List Number"]._element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(1)
    _list_num["id"] = num.numId


def steps(items):
    for it in items:
        par = p(it, "List Number")
        numPr = par._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0
        numPr.get_or_add_numId().val = _list_num["id"]


def code(*lines):
    for i, line in enumerate(lines):
        par = doc.add_paragraph()
        shade(par._p, "EEF1F5")  # before spacing/indent so w:shd lands in schema order
        pf = par.paragraph_format
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(6) if i == len(lines) - 1 else Pt(0)
        pf.left_indent = Cm(0.4)
        pf.line_spacing = 1.0
        r = par.add_run(line or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)


def note(text, fill="FFF4D6", label="Note"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    shade(cell._tc, fill)
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(label + ": ")
    r.bold = True
    rich(par, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c._tc, "1D4E89")
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            rich(cells[i].paragraphs[0], txt)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                for r in par.runs:
                    if r.font.name != "Consolas":
                        r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------- title
title = doc.add_paragraph()
r = title.add_run("KRDP Remote Desktop on CachyOS Plasma Wayland")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = ACCENT
sub = doc.add_paragraph()
r = sub.add_run("Setup, fixes applied, and the Plasma 6.8 upgrade")
r.font.size = Pt(12)
r.font.color.rgb = MUTED
meta = doc.add_paragraph()
r = meta.add_run("Host: cachyos-taichi (192.168.0.101)  ·  Written 15 October 2026, updated 16 October 2026 (Plasma 6.8)")
r.font.size = Pt(9.5)
r.font.color.rgb = MUTED

# ---------------------------------------------------------------- summary
h("1. Summary")
p("Remote Desktop into the live Plasma Wayland session (via KDE's KRDP server) failed for every "
  "client. Three problems were fixed on Plasma 6.7.5. After the upgrade to Plasma 6.8.0, a fourth "
  "problem (a black screen) appeared and was fixed by installing krdp 6.8.1. A Windows PC now logs "
  "in as user `hs` with the normal Linux password and sees the desktop.")
table(["#", "Problem", "Fix applied"], [
    ["1", "FreeRDP 3.32 turned on Extended Security (HYBRID_EX) by default, so KRDP was forced into "
          "an NLA login it never configures (“Could not find user in SAM database”).",
     "Rebuilt krdp 6.7.5 with KDE's upstream patch (disables ExtSecurity)."],
    ["2", "Windows sends the username as `PCNAME\\hs` or `.\\hs`. KRDP 6.7.5 compares it to `hs` "
          "exactly, so the password was never checked.",
     "Added a patch that strips the `DOMAIN\\` prefix (backport of upstream master)."],
    ["3", "The auto-generated TLS certificate was valid for only 1 day, had no hostname, and is "
          "regenerated whenever the Remote Desktop settings page opens.",
     "Created a 10-year certificate for cachyos-taichi / 192.168.0.101 and set "
     "`AutogenerateCertificates=false`."],
    ["4", "Plasma 6.8.0: black screen after login. KRDP's new video flow control waited for frame "
          "acknowledgements that Windows had suspended, so the stream froze after 6 frames.",
     "Built and installed krdp 6.8.1 (KDE release commit, includes the fix for BUG 520260)."],
], [0.8, 8.6, 7.2])

# ---------------------------------------------------------------- system
h("2. Current system")
table(["Component", "Version", "Notes"], [
    ["Plasma / plasma-desktop", "6.8.0", "Wayland session (upgraded from 6.7.5 on 16 Oct 2026)"],
    ["krdp", "6.8.1-1 (local build)", "CachyOS ships 6.8.0, which gives a black screen with Windows"],
    ["GPU", "NVIDIA RTX 3090 Ti", "No VA-API H.264 encode; KRDP uses software x264"],
    ["freerdp", "3.32.2-1.1", "Introduced the ExtSecurity default change"],
    ["Authentication", "System user (PAM)", "User `hs`, normal Linux password"],
    ["Port", "TCP 3389", "Service: `app-org.kde.krdpserver.service` (user unit)"],
], [4.5, 4.4, 7.7])

# ---------------------------------------------------------------- root causes
h("3. What was wrong, in detail")
h("3.1 FreeRDP 3.32 Extended Security default", 2)
p("FreeRDP commit a1fad13 (“add ExtSecurity to all clients and server”, September 2026) changed "
  "`FreeRDP_ExtSecurity` from FALSE to TRUE by default. KRDP 6.7.5 is designed to use TLS with a "
  "PAM password check, and disables `NlaSecurity`, but it did not know about the new setting. "
  "Clients therefore negotiated PROTOCOL_HYBRID_EX (a form of NLA/CredSSP), and FreeRDP looked "
  "for the user in an NTLM “SAM” database that KRDP never creates. Every login failed with:")
code("[ERROR][com.winpr.sspi.NTLM] - [ntlm_fetch_ntlm_v2_hash]: Could not find user in SAM database",
     "AcceptSecurityContext status SEC_E_NO_CREDENTIALS [0x8009030e]")
p("KDE fixed this in commit 573e40a on the Plasma/6.7 branch (“Fix authentication with FreeRDP 3.32”, "
  "KDE bug 526514), but no 6.7.x release contained it yet.")

h("3.2 Windows username format", 2)
p("After fix 1, Windows reached the server but was rejected without a PAM check. KRDP 6.7.5 only "
  "checks the password if the username equals the local login name (`hs`) exactly. Remote Desktop "
  "Connection typically sends `PCNAME\\hs`, `.\\hs` or the Windows account name. Upstream master "
  "commit 095f6ea strips the prefix; that change was backported here.")

h("3.3 One-day TLS certificate", 2)
p("The Remote Desktop settings page (6.7.5) runs `openssl req … -days 1` every time it opens while "
  "`AutogenerateCertificates` is true (the default). That produced a certificate expiring the next "
  "day, with no hostname, and Windows showed “The remote computer could not be authenticated”. "
  "The certificate was replaced and auto-generation disabled.")

# ---------------------------------------------------------------- changes
h("3.4 Black screen on Plasma 6.8.0", 2)
p("After the 6.8.0 upgrade, login worked but Windows showed a black screen. x264 encoded one keyframe "
  "with real desktop content and five small frames, then stopped. krdp 6.8.0 rewrote its video flow "
  "control: it limits frames “in flight” until the client acknowledges them. Windows can suspend "
  "acknowledgements (`SUSPEND_FRAME_ACKNOWLEDGEMENT`). 6.8.0 kept waiting, the window filled, and "
  "the encoder paused forever. KDE fixed this in krdp 6.8.1 (BUG 520260).")

h("4. Changes made on this machine")
table(["Change", "Location", "How to undo"], [
    ["Package krdp 6.8.1-1 (local build of KDE's release commit 5813706)",
     "Build files: `scripts/rdp/krdp-build/PKGBUILD` (cachyos-setup repo)",
     "`sudo pacman -S krdp` reinstalls the CachyOS build (6.8.0 gives a black screen)"],
    ["New TLS certificate (10 years)",
     "`~/.local/share/krdpserver/krdp.crt` and `krdp.key`",
     "Originals saved as `krdp.crt.bak-20261015` / `krdp.key.bak-20261015`"],
    ["Certificate auto-generation off",
     "`~/.config/krdpserverrc`: `AutogenerateCertificates=false`",
     "`kwriteconfig6 --file krdpserverrc --group General --key AutogenerateCertificates true`"],
], [4.2, 6.2, 6.2])
p("On Plasma 6.7.5 the package was krdp 6.7.5-3 with two patches (ExtSecurity and `DOMAIN\\user`). "
  "Both fixes are included in 6.8.1, so the current build has no local patches.")
note("The CachyOS package `6.8.1-1.1`, when released, sorts higher than the local `6.8.1-1`, so a "
     "normal `pacman -Syu` will replace it automatically. **Do not** add krdp to IgnorePkg.",
     fill="E3F0FF", label="Updates")

# ---------------------------------------------------------------- connecting
h("5. Connecting from Windows")
new_list()
steps([
    "Open **Remote Desktop Connection** (`mstsc`) and click **Show Options**.",
    "Computer: `192.168.0.101` (or `cachyos-taichi` if your network resolves it).",
    "User name: `hs`. Tick **Allow me to save credentials**.",
    "If Windows has saved credentials for this PC that may be old, click **Delete** first.",
    "Click **Connect** and enter the **Linux** password for `hs` (not the Windows/Microsoft password).",
    "At “The remote computer could not be authenticated due to problems with its security certificate”, "
    "tick **Don't ask me again for connections to this computer** and click **Yes**. This is expected "
    "with a self-signed certificate and will not reappear unless the certificate changes.",
])
note("The Linux side must be logged in to Plasma with the screen available; KRDP shares the "
     "existing session and does not create a new one. A KDE prompt may ask to allow remote control "
     "the first time.", label="Note")
note("CachyOS locks the account after 3 failed logins in 15 minutes (pam_faillock), and the lock "
     "also blocks local login for about 10 minutes. Clear it with `faillock --user hs --reset`.",
     fill="FDE7E7", label="Warning")
p("Avoid non-ASCII characters (£, €, accented letters) in the password: KRDP reads RDP passwords "
  "as Latin-1, so such characters may not match. This is still true in 6.8.0.")

# ---------------------------------------------------------------- 6.8
h("6. Plasma 6.8 status")
p("The upgrade to Plasma 6.8.0 happened on 16 October 2026. Status of each known issue with krdp 6.8.1:")
table(["Issue", "Status in 6.8.1", "What to do"], [
    ["ExtSecurity / “SAM database” error", "**Fixed** (since 6.8.0).", "Nothing."],
    ["Windows `DOMAIN\\user` username", "**Fixed** in 6.8.1.", "Nothing."],
    ["Black screen / frozen picture with Windows", "**Fixed** in 6.8.1 (broken in 6.8.0).",
     "Keep krdp at 6.8.1 or later. Never go back to 6.8.0."],
    ["One-day certificate", "**Partly**: 6.8 only generates one if no certificate is configured, "
     "but it is still 1 day.",
     "Keep `AutogenerateCertificates=false` and the 10-year certificate. Check after upgrade."],
    ["Non-ASCII passwords", "**No**: still Latin-1.", "Keep the password ASCII."],
    ["“Other Users” with empty password", "**No**: any such entry ends the check early.",
     "Remove the leftover diagnostic `test` user in System Settings → Remote Desktop."],
], [4.4, 5.6, 6.6])

h("6.1 Upgrade checklist", 2)
new_list()
steps([
    "Upgrade normally: `sudo pacman -Syu`. Confirm krdp is 6.8.1 or newer:",
])
code("pacman -Q plasma-desktop krdp freerdp")
steps([
    "Restart the server (or log out and back in):",
])
code("systemctl --user restart app-org.kde.krdpserver.service",
     "systemctl --user is-active app-org.kde.krdpserver.service",
     "ss -ltn | grep 3389")
steps([
    "Confirm the certificate was not replaced (should say 2036):",
])
code("openssl x509 -in ~/.local/share/krdpserver/krdp.crt -noout -subject -enddate",
     "kreadconfig6 --file krdpserverrc --group General --key AutogenerateCertificates")
steps([
    "Connect from Windows as `hs`. If it fails, read the log (section 8).",
])
note("6.8 also adds a separate NLA mode for “Other Users” (used when the System User option is off). "
     "This setup uses the System User (PAM) mode, which is unaffected.", label="Note")

# ---------------------------------------------------------------- rebuild
h("7. Rebuilding with patches (only if needed)")
p("Use this when a KDE fix exists upstream but CachyOS hasn't packaged it yet (as with 6.8.1). "
  "Build tools: `sudo pacman -S --needed base-devel cmake extra-cmake-modules plasma-wayland-protocols git`.")
new_list()
steps([
    "In `scripts/rdp/krdp-build/PKGBUILD` (cachyos-setup repo), set `pkgver` and `_commit` to the KDE release commit "
    "(see `git ls-remote https://invent.kde.org/plasma/krdp.git refs/heads/Plasma/6.8`). Use `pkgrel=1` "
    "so the CachyOS package of the same version (`-1.1`) replaces it automatically later.",
    "Build tools must be installed (not as dependencies): "
    "`sudo pacman -S --needed cmake extra-cmake-modules plasma-wayland-protocols`.",
    "Always keep `-DCMAKE_BUILD_TYPE=Release` (debug builds hit an assertion and crash on connect).",
    "Build and install:",
])
code("cd cachyos-setup/scripts/rdp/krdp-build",
     "rm -rf src pkg; rm -f *.pkg.tar.zst",
     "makepkg -f",
     "sudo pacman -U krdp-*.pkg.tar.zst      # run in a normal terminal",
     "systemctl --user restart app-org.kde.krdpserver.service",
     "pacman -Q krdp                          # confirm the new version")
note("A krdpserver binary run from the build folder crashes as soon as a client connects. That is "
     "expected: KWin only gives the screencast protocol to `/usr/bin/krdpserver`. Test only after installing.",
     label="Testing")

# ---------------------------------------------------------------- troubleshooting
h("8. Troubleshooting")
p("Live log while connecting:")
code("journalctl --user -u app-org.kde.krdpserver.service -f")
table(["Log message / symptom", "Meaning", "Fix"], [
    ["`Could not find user in SAM database` / `SEC_E_NO_CREDENTIALS`",
     "NLA/ExtSecurity selected; KRDP is not handling the login.",
     "krdp lacks the ExtSecurity fix. Install 6.8.0+ or rebuild with the patch."],
    ["`New client connected` then `PostConnect … failed` with no `pam_` line",
     "Username did not match `hs`.",
     "Enter `hs` exactly in mstsc, or rebuild with the strip patch."],
    ["`pam_unix(login:auth): authentication failure`",
     "Wrong password received.",
     "Use the Linux password; clear saved credentials in mstsc; avoid non-ASCII."],
    ["`Authentication failed for user … (password empty: true)` (patched build only)",
     "Client sent no password.", "Tick “Allow me to save credentials” and enter the password."],
    ["Windows: “remote computer could not be authenticated”",
     "Self-signed certificate warning.",
     "Accept once with “Don't ask me again”. If it keeps returning, the certificate changed (section 6.1)."],
    ["Login works locally but not after many tries", "faillock lockout.", "`faillock --user hs --reset`"],
    ["Black screen after a successful login", "Video stream stalled (krdp 6.8.0 flow-control bug).",
     "Use krdp 6.8.1+. To confirm a stall: server CPU stays at 0 while connected, and x264 reports only a few frames when the session ends."],
    ["Port 3389 not listening", "Service stopped or crashed.",
     "`systemctl --user restart app-org.kde.krdpserver.service`; check `coredumpctl list krdpserver`."],
], [5.4, 4.6, 6.6])

h("9. Regenerating the certificate")
p("If the certificate is lost or the IP address changes:")
code("cd ~/.local/share/krdpserver",
     "openssl req -x509 -newkey rsa:2048 -nodes -days 3650 \\",
     "  -keyout krdp.key -out krdp.crt -subj \"/CN=$(hostname)\" \\",
     "  -addext \"subjectAltName=DNS:$(hostname),IP:192.168.0.101\" \\",
     "  -addext \"extendedKeyUsage=serverAuth\"",
     "chmod 600 krdp.key",
     "systemctl --user restart app-org.kde.krdpserver.service")
p("Windows will show the certificate warning once more after this.")

h("10. References")
bullets([
    "KDE bug 526514: KRDP authentication broken with FreeRDP 3.32.",
    "KRDP commit 573e40a (Plasma/6.7): Fix authentication with FreeRDP 3.32.",
    "KRDP commit 095f6ea (master): Fix Windows RDP client (mstsc) connection and authentication.",
    "FreeRDP commit a1fad13: [core,security] add ExtSecurity to all clients and server.",
    "KDE bug 520260 / KRDP commit 472584e (Plasma/6.8): Stop waiting for acknowledgements the client suspended.",
    "Source: https://invent.kde.org/plasma/krdp",
])

# footer page numbers
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    run._r.append(el)
run.font.size = Pt(9)
run.font.color.rgb = MUTED

zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")

doc.save(OUT)
print(OUT)
